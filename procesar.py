#!/usr/bin/env python3
"""
Procesador de Calificaciones Google Classroom + Rúbricas
=========================================================

Lee el archivo CSV/Excel exportado de Google Classroom y genera un reporte
Excel que desglosa los criterios de rúbrica por actividad.

Flujo de uso:
  1. python procesar.py plantilla  <classroom.csv>
     → Genera plantilla_rubricas.xlsx para configurar rúbricas

  2. python procesar.py datos  <classroom.csv>  <plantilla_rubricas.xlsx>
     → Agrega hojas de ingreso de punteos por criterio

  3. python procesar.py reporte  <classroom.csv>  <plantilla_rubricas.xlsx>
     → Genera reporte_final.xlsx con todo el desglose
"""

import sys
import os
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paleta de colores ────────────────────────────────────────────────────────
AZUL_OSCURO  = "1F4E79"
AZUL_MEDIO   = "2E75B6"
AZUL_CLARO   = "BDD7EE"
VERDE_OSCURO = "375623"
VERDE_MEDIO  = "548235"
VERDE_CLARO  = "E2EFDA"
GRIS_CLARO   = "F2F2F2"
GRIS_MEDIO   = "D9D9D9"
AMARILLO     = "FFF2CC"
NARANJA      = "FCE4D6"
BLANCO       = "FFFFFF"


def _celda(ws, fila, col, valor=None, negrita=False, color_fuente=BLANCO,
           bg=None, alinear="center", wrap=False, tamaño=10, italica=False):
    """Aplica formato a una celda y opcionalmente asigna valor."""
    cell = ws.cell(fila, col, valor)
    cell.font = Font(bold=negrita, italic=italica, color=color_fuente, size=tamaño)
    cell.alignment = Alignment(horizontal=alinear, vertical="center", wrap_text=wrap)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    return cell


def _borde_fino(ws, fila_ini, fila_fin, col_ini, col_fin):
    thin = Side(style="thin", color="AAAAAA")
    b = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(fila_ini, fila_fin + 1):
        for c in range(col_ini, col_fin + 1):
            ws.cell(r, c).border = b


# ── Lector del formato Classroom ────────────────────────────────────────────

def leer_classroom(ruta):
    """
    Lee el CSV/Excel exportado de Google Classroom.
    Estructura esperada:
      Fila 0: Nombre del curso | ...fechas de actividades...
      Fila 1: Sección          | ...nombres de actividades...
      Fila 2: Link Classroom   | ...puntos máximos...
      Fila 3: (vacía)
      Fila 4: Promedio de la clase
      Fila 5+: Estudiantes (Apellido | Nombre | Email | % | punteos...)
    """
    ext = os.path.splitext(ruta)[1].lower()
    if ext in (".xlsx", ".xls"):
        df = pd.read_excel(ruta, header=None, dtype=str)
    else:
        df = pd.read_csv(ruta, header=None, dtype=str)

    curso   = str(df.iloc[0, 0]).strip()
    seccion = str(df.iloc[1, 0]).strip()

    # Actividades: desde columna índice 4 en adelante
    actividades = []
    for ci in range(4, df.shape[1]):
        nombre = str(df.iloc[1, ci]).strip()
        fecha  = str(df.iloc[0, ci]).strip()
        max_p  = df.iloc[2, ci]
        if nombre and nombre.lower() not in ("nan", "none", ""):
            try:
                max_p = float(str(max_p).replace(",", "."))
            except (ValueError, TypeError):
                max_p = 0.0
            actividades.append({
                "nombre":     nombre,
                "fecha":      fecha if fecha.lower() not in ("nan", "none", "") else "",
                "max_puntos": max_p,
                "col_idx":    ci,
            })

    # Estudiantes: desde fila índice 5
    estudiantes = []
    for ri in range(5, df.shape[0]):
        apellido = str(df.iloc[ri, 0]).strip()
        if not apellido or apellido.lower() in ("nan", "none", "promedio de la clase"):
            continue
        nombre_e = str(df.iloc[ri, 1]).strip()
        email    = str(df.iloc[ri, 2]).strip()
        pct      = df.iloc[ri, 3]

        est = {"apellido": apellido, "nombre": nombre_e,
               "email": email,      "porcentaje": pct}

        for act in actividades:
            raw = df.iloc[ri, act["col_idx"]]
            try:
                est[act["nombre"]] = float(str(raw).replace(",", "."))
            except (ValueError, TypeError):
                est[act["nombre"]] = None

        estudiantes.append(est)

    return {
        "curso":        curso,
        "seccion":      seccion,
        "actividades":  actividades,
        "estudiantes":  estudiantes,
    }


# ── PASO 1: Generar plantilla ────────────────────────────────────────────────

def generar_plantilla(ruta_classroom, salida="plantilla_rubricas.xlsx"):
    data = leer_classroom(ruta_classroom)
    wb   = openpyxl.Workbook()

    # ── Hoja Configuracion ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Configuracion"

    # Título
    ws.merge_cells("A1:L1")
    _celda(ws, 1, 1,
           f"CONFIGURACIÓN DE RÚBRICAS — {data['curso']}  |  {data['seccion']}",
           negrita=True, bg=AZUL_OSCURO, tamaño=13)
    ws.row_dimensions[1].height = 30

    # Instrucciones
    instrucciones = [
        "INSTRUCCIONES:",
        "1. En la columna 'Tiene Rúbrica' escriba  SÍ  o  NO  para cada actividad.",
        "2. Para las actividades con rúbrica, escriba el nombre de cada criterio y sus puntos máximos.",
        "   Ejemplo: Criterio 1 = Puntualidad, Máx C1 = 5  |  Criterio 2 = Contenido, Máx C2 = 5  ...",
        "3. Guarde este archivo.",
        "4. Ejecute:  python procesar.py datos  <classroom.csv>  plantilla_rubricas.xlsx",
        "   Esto generará una hoja por actividad con rúbrica para ingresar los punteos.",
    ]
    for i, txt in enumerate(instrucciones, 2):
        ws.merge_cells(f"A{i}:L{i}")
        color = "000000" if i > 2 else AZUL_OSCURO
        bold  = i == 2
        _celda(ws, i, 1, txt, negrita=bold, color_fuente=color,
               bg=AZUL_CLARO if i == 2 else GRIS_CLARO,
               alinear="left", wrap=True, tamaño=9, italica=(i > 2))
        ws.row_dimensions[i].height = 14

    # Encabezados tabla
    HR = 2 + len(instrucciones)
    headers = [
        "Actividad", "Fecha", "Máx. Pts", "¿Tiene Rúbrica?",
        "Criterio 1", "Máx C1",
        "Criterio 2", "Máx C2",
        "Criterio 3", "Máx C3",
        "Criterio 4", "Máx C4",
    ]
    for ci, h in enumerate(headers, 1):
        _celda(ws, HR, ci, h, negrita=True, bg=AZUL_MEDIO, wrap=True)
    ws.row_dimensions[HR].height = 35

    # Anchos
    anchos = [32, 10, 9, 16, 20, 8, 20, 8, 20, 8, 20, 8]
    for ci, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # Filas de actividades
    for ri, act in enumerate(data["actividades"]):
        r  = HR + 1 + ri
        bg = AZUL_CLARO if ri % 2 == 0 else BLANCO
        _celda(ws, r, 1, act["nombre"],     bg=bg, alinear="left", color_fuente="000000")
        _celda(ws, r, 2, act["fecha"],      bg=bg, color_fuente="000000")
        _celda(ws, r, 3, act["max_puntos"], bg=bg, color_fuente="000000")
        _celda(ws, r, 4, "NO",              bg=GRIS_MEDIO, color_fuente="000000")
        for ci in range(5, 13):
            _celda(ws, r, ci, None, bg=AMARILLO, color_fuente="000000")

    _borde_fino(ws, HR, HR + len(data["actividades"]), 1, 12)
    wb.save(salida)

    print(f"\n✅  Plantilla generada: {salida}")
    print("\n📋  Pasos siguientes:")
    print(f"   1. Abra '{salida}' en Excel o Google Sheets.")
    print(f"   2. En la columna 'Tiene Rúbrica' escriba SÍ o NO.")
    print(f"   3. Para las actividades con rúbrica, llene los criterios y sus puntos máximos.")
    print(f"   4. Ejecute:  python procesar.py datos {ruta_classroom} {salida}")


# ── PASO 2: Agregar hojas de ingreso de punteos ─────────────────────────────

def _leer_config(ws_config):
    """Extrae la configuración de rúbricas de la hoja Configuracion."""
    header_row = None
    for row in ws_config.iter_rows():
        for cell in row:
            if str(cell.value).strip() == "Actividad":
                header_row = cell.row
                break
        if header_row:
            break
    if not header_row:
        print("❌  No se encontró la tabla en la hoja 'Configuracion'.")
        sys.exit(1)

    config = {}
    for row in ws_config.iter_rows(min_row=header_row + 1, values_only=True):
        if not row[0]:
            continue
        nombre = str(row[0]).strip()
        rubrica = str(row[3]).strip().upper() if row[3] else "NO"
        criterios = []
        if rubrica in ("SÍ", "SI", "S"):
            for i in range(4):
                nc = row[4 + i * 2] if len(row) > 4 + i * 2 else None
                mx = row[5 + i * 2] if len(row) > 5 + i * 2 else None
                if nc and str(nc).strip() not in ("", "nan", "None"):
                    try:
                        mx = float(str(mx).replace(",", ".")) if mx else 0.0
                    except (ValueError, TypeError):
                        mx = 0.0
                    criterios.append({"nombre": str(nc).strip(), "max": mx})
        config[nombre] = {
            "tiene_rubrica": rubrica in ("SÍ", "SI", "S"),
            "criterios": criterios,
        }
    return config


def generar_hojas_rubrica(ruta_classroom, ruta_plantilla):
    data   = leer_classroom(ruta_classroom)
    wb     = openpyxl.load_workbook(ruta_plantilla)
    config = _leer_config(wb["Configuracion"])

    count = 0
    for act in data["actividades"]:
        cfg = config.get(act["nombre"], {"tiene_rubrica": False, "criterios": []})
        if not cfg["tiene_rubrica"] or not cfg["criterios"]:
            continue

        sheet_name = act["nombre"][:28].replace("/", "-")
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)

        criterios  = cfg["criterios"]
        total_cols = 3 + len(criterios) + 1   # Apellido | Nombre | Email | criterios... | TOTAL

        # Título
        ws.merge_cells(f"A1:{get_column_letter(total_cols)}1")
        _celda(ws, 1, 1,
               f"RÚBRICA: {act['nombre']}   |   Máx: {act['max_puntos']} pts   |   "
               f"{data['curso']}  {data['seccion']}",
               negrita=True, bg=VERDE_OSCURO, tamaño=11)
        ws.row_dimensions[1].height = 26

        # Fila de máximos
        _celda(ws, 2, 1, None, bg=GRIS_CLARO)
        _celda(ws, 2, 2, None, bg=GRIS_CLARO)
        _celda(ws, 2, 3, None, bg=GRIS_CLARO)
        for ci, crit in enumerate(criterios):
            _celda(ws, 2, 4 + ci, f"Máx: {crit['max']}",
                   italica=True, bg=GRIS_CLARO, color_fuente="555555", tamaño=8)
        _celda(ws, 2, 4 + len(criterios), f"Máx Total: {act['max_puntos']}",
               negrita=True, italica=True, bg=GRIS_CLARO, color_fuente="555555", tamaño=8)
        ws.row_dimensions[2].height = 14

        # Encabezados
        encabezados = ["Apellido", "Nombre", "Email"] + \
                      [c["nombre"] for c in criterios] + ["TOTAL"]
        for ci, h in enumerate(encabezados, 1):
            _celda(ws, 3, ci, h, negrita=True, bg=VERDE_MEDIO, wrap=True)
        ws.row_dimensions[3].height = 35

        # Anchos
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 30
        for ci in range(len(criterios) + 1):
            ws.column_dimensions[get_column_letter(4 + ci)].width = 14

        # Filas de estudiantes
        for ri, est in enumerate(data["estudiantes"]):
            r  = 4 + ri
            bg = VERDE_CLARO if ri % 2 == 0 else BLANCO
            _celda(ws, r, 1, est["apellido"], bg=bg, alinear="left", color_fuente="000000")
            _celda(ws, r, 2, est["nombre"],   bg=bg, alinear="left", color_fuente="000000")
            _celda(ws, r, 3, est["email"],    bg=bg, alinear="left", color_fuente="000000")
            for ci in range(len(criterios)):
                _celda(ws, r, 4 + ci, None, bg=AMARILLO, color_fuente="000000")
            # Fórmula TOTAL = suma de criterios
            ini = get_column_letter(4)
            fin = get_column_letter(3 + len(criterios))
            ws.cell(r, 4 + len(criterios)).value = f"=SUM({ini}{r}:{fin}{r})"
            ws.cell(r, 4 + len(criterios)).font  = Font(bold=True, size=10)
            ws.cell(r, 4 + len(criterios)).fill  = PatternFill("solid", fgColor=VERDE_CLARO)
            ws.cell(r, 4 + len(criterios)).alignment = Alignment(horizontal="center", vertical="center")

        _borde_fino(ws, 3, 3 + len(data["estudiantes"]), 1, total_cols)
        count += 1

    wb.save(ruta_plantilla)
    print(f"\n✅  {count} hoja(s) de rúbrica generadas en: {ruta_plantilla}")
    print("\n📋  Pasos siguientes:")
    print(f"   1. Abra '{ruta_plantilla}'.")
    print(f"   2. En cada hoja amarilla, ingrese el punteo de cada criterio por estudiante.")
    print(f"   3. Ejecute:  python procesar.py reporte {ruta_classroom} {ruta_plantilla}")


# ── PASO 3: Reporte final ────────────────────────────────────────────────────

def generar_reporte(ruta_classroom, ruta_plantilla, salida="reporte_final.xlsx"):
    data       = leer_classroom(ruta_classroom)
    wb_src     = openpyxl.load_workbook(ruta_plantilla, data_only=True)
    config     = _leer_config(wb_src["Configuracion"])

    # Leer punteos de rúbricas de cada hoja
    rubrica_data = {}   # {nombre_act: {email: {criterio: pts, "TOTAL": pts}}}
    for act in data["actividades"]:
        cfg = config.get(act["nombre"], {"tiene_rubrica": False, "criterios": []})
        if not cfg["tiene_rubrica"] or not cfg["criterios"]:
            continue
        sheet_name = act["nombre"][:28].replace("/", "-")
        if sheet_name not in wb_src.sheetnames:
            continue
        ws_act = wb_src[sheet_name]
        criterios = cfg["criterios"]
        rubrica_data[act["nombre"]] = {}
        for row in ws_act.iter_rows(min_row=4, values_only=True):
            if not row[2]:
                continue
            email = str(row[2]).strip()
            scores = {}
            for ci, crit in enumerate(criterios):
                val = row[3 + ci]
                try:
                    scores[crit["nombre"]] = float(val) if val is not None else 0.0
                except (ValueError, TypeError):
                    scores[crit["nombre"]] = 0.0
            total_val = row[3 + len(criterios)]
            try:
                scores["TOTAL"] = float(total_val) if total_val is not None else sum(scores.values())
            except (ValueError, TypeError):
                scores["TOTAL"] = sum(scores.values())
            rubrica_data[act["nombre"]][email] = scores

    # ── Construir Excel de reporte ──────────────────────────────────────────
    wb_out = openpyxl.Workbook()
    ws     = wb_out.active
    ws.title = "Reporte Final"

    # Calcular número total de columnas
    act_cols = []   # lista de (act, n_cols)
    for act in data["actividades"]:
        cfg = config.get(act["nombre"], {"tiene_rubrica": False, "criterios": []})
        if cfg["tiene_rubrica"] and cfg["criterios"]:
            act_cols.append((act, cfg, len(cfg["criterios"]) + 1))
        else:
            act_cols.append((act, cfg, 1))
    total_cols = 4 + sum(nc for _, _, nc in act_cols)

    # ── Fila 1: Título principal ────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(total_cols)}1")
    _celda(ws, 1, 1,
           f"REPORTE DE CALIFICACIONES  —  {data['curso']}    {data['seccion']}",
           negrita=True, bg=AZUL_OSCURO, tamaño=14)
    ws.row_dimensions[1].height = 34

    # ── Fila 2: Encabezados fijos + bloques de actividad ───────────────────
    for c, h in enumerate(["Apellido", "Nombre", "Email", "%"], 1):
        ws.merge_cells(f"{get_column_letter(c)}2:{get_column_letter(c)}3")
        _celda(ws, 2, c, h, negrita=True, bg=AZUL_OSCURO)

    col = 5
    for act, cfg, nc in act_cols:
        tiene_rubrica = cfg["tiene_rubrica"] and cfg["criterios"]
        bg_act = VERDE_OSCURO if tiene_rubrica else AZUL_MEDIO
        if nc > 1:
            ws.merge_cells(f"{get_column_letter(col)}2:{get_column_letter(col+nc-1)}2")
        _celda(ws, 2, col,
               f"{act['nombre']}\n{act['fecha']}",
               negrita=True, bg=bg_act, wrap=True, tamaño=9)
        col += nc
    ws.row_dimensions[2].height = 50

    # ── Fila 3: Sub-criterios ───────────────────────────────────────────────
    col = 5
    for act, cfg, nc in act_cols:
        tiene_rubrica = cfg["tiene_rubrica"] and cfg["criterios"]
        if tiene_rubrica:
            for crit in cfg["criterios"]:
                _celda(ws, 3, col, crit["nombre"],
                       negrita=True, bg=VERDE_MEDIO, wrap=True, tamaño=8)
                col += 1
            _celda(ws, 3, col, "TOTAL", negrita=True, bg=VERDE_OSCURO, tamaño=9)
            col += 1
        else:
            _celda(ws, 3, col, "Punteo", negrita=True, bg=AZUL_MEDIO, tamaño=9)
            col += 1
    ws.row_dimensions[3].height = 38

    # ── Fila 4: Puntos máximos ──────────────────────────────────────────────
    _celda(ws, 4, 1, "PUNTOS MÁXIMOS", italica=True,
           bg=GRIS_CLARO, alinear="right", color_fuente="555555", tamaño=8)
    ws.merge_cells("A4:D4")
    col = 5
    for act, cfg, nc in act_cols:
        tiene_rubrica = cfg["tiene_rubrica"] and cfg["criterios"]
        if tiene_rubrica:
            for crit in cfg["criterios"]:
                _celda(ws, 4, col, crit["max"],
                       italica=True, bg=GRIS_CLARO, color_fuente="555555", tamaño=8)
                col += 1
            _celda(ws, 4, col, act["max_puntos"],
                   negrita=True, italica=True, bg=GRIS_CLARO, color_fuente="555555", tamaño=8)
            col += 1
        else:
            _celda(ws, 4, col, act["max_puntos"],
                   italica=True, bg=GRIS_CLARO, color_fuente="555555", tamaño=8)
            col += 1
    ws.row_dimensions[4].height = 16

    # ── Anchos de columnas ──────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 8
    col = 5
    for act, cfg, nc in act_cols:
        tiene_rubrica = cfg["tiene_rubrica"] and cfg["criterios"]
        if tiene_rubrica:
            for _ in cfg["criterios"]:
                ws.column_dimensions[get_column_letter(col)].width = 14
                col += 1
            ws.column_dimensions[get_column_letter(col)].width = 9
            col += 1
        else:
            ws.column_dimensions[get_column_letter(col)].width = 16
            col += 1

    # ── Filas de estudiantes ────────────────────────────────────────────────
    for ri, est in enumerate(data["estudiantes"]):
        r  = 5 + ri
        bg = AZUL_CLARO if ri % 2 == 0 else BLANCO

        _celda(ws, r, 1, est["apellido"],   bg=bg, alinear="left", color_fuente="000000")
        _celda(ws, r, 2, est["nombre"],     bg=bg, alinear="left", color_fuente="000000")
        _celda(ws, r, 3, est["email"],      bg=bg, alinear="left", color_fuente="000000")
        pct_cell = ws.cell(r, 4)
        try:
            pct_val = float(str(est["porcentaje"]).replace("%", "").replace(",", "."))
            pct_cell.value          = pct_val / 100
            pct_cell.number_format  = "0.00%"
        except (ValueError, TypeError):
            pct_cell.value = est["porcentaje"]
        pct_cell.fill      = PatternFill("solid", fgColor=bg)
        pct_cell.alignment = Alignment(horizontal="center", vertical="center")

        col = 5
        for act, cfg, nc in act_cols:
            tiene_rubrica = cfg["tiene_rubrica"] and cfg["criterios"]
            if tiene_rubrica:
                email  = est["email"]
                rb_est = rubrica_data.get(act["nombre"], {}).get(email, {})
                for crit in cfg["criterios"]:
                    val = rb_est.get(crit["nombre"], "")
                    cell = ws.cell(r, col, val)
                    cell.fill      = PatternFill("solid", fgColor=bg)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    col += 1
                total_val = rb_est.get("TOTAL", est.get(act["nombre"], ""))
                cell = ws.cell(r, col, total_val)
                cell.font      = Font(bold=True, size=10)
                cell.fill      = PatternFill("solid", fgColor=VERDE_CLARO)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                col += 1
            else:
                val  = est.get(act["nombre"], "")
                cell = ws.cell(r, col, val)
                cell.fill      = PatternFill("solid", fgColor=bg)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                col += 1

    _borde_fino(ws, 2, 4 + len(data["estudiantes"]), 1, total_cols)
    ws.freeze_panes = "E5"

    wb_out.save(salida)

    print(f"\n✅  Reporte final generado: {salida}")
    print(f"\n📊  Resumen:")
    print(f"    Curso:       {data['curso']}  —  {data['seccion']}")
    print(f"    Estudiantes: {len(data['estudiantes'])}")
    print(f"    Actividades: {len(data['actividades'])}")
    for act, cfg, nc in act_cols:
        if cfg["tiene_rubrica"] and cfg["criterios"]:
            criterios_str = ", ".join(
                f"{c['nombre']} (/{c['max']})" for c in cfg["criterios"]
            )
            print(f"      ✦ {act['nombre']} — Rúbrica: {criterios_str}")
        else:
            print(f"      ○ {act['nombre']} — Sin rúbrica")


# ── Punto de entrada ─────────────────────────────────────────────────────────

def main():
    ayuda = """
Uso:
  python procesar.py plantilla  <classroom.csv>
      Genera 'plantilla_rubricas.xlsx' para configurar criterios de rúbrica.

  python procesar.py datos  <classroom.csv>  <plantilla_rubricas.xlsx>
      Agrega hojas de ingreso de punteos por criterio (celdas amarillas).

  python procesar.py reporte  <classroom.csv>  <plantilla_rubricas.xlsx>  [salida.xlsx]
      Genera el reporte final combinando Classroom + rúbricas.
"""
    if len(sys.argv) < 3:
        print(ayuda)
        sys.exit(1)

    comando = sys.argv[1].lower()

    if comando == "plantilla":
        generar_plantilla(sys.argv[2])

    elif comando == "datos":
        if len(sys.argv) < 4:
            print("❌  Falta la plantilla: python procesar.py datos <classroom.csv> <plantilla.xlsx>")
            sys.exit(1)
        generar_hojas_rubrica(sys.argv[2], sys.argv[3])

    elif comando == "reporte":
        if len(sys.argv) < 4:
            print("❌  Falta la plantilla: python procesar.py reporte <classroom.csv> <plantilla.xlsx>")
            sys.exit(1)
        salida = sys.argv[4] if len(sys.argv) > 4 else "reporte_final.xlsx"
        generar_reporte(sys.argv[2], sys.argv[3], salida)

    else:
        print(f"❌  Comando desconocido: '{comando}'")
        print("    Comandos válidos: plantilla, datos, reporte")
        sys.exit(1)


if __name__ == "__main__":
    main()
